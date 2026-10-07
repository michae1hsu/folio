"""Local PDF conversion and preview pages, including Office supplements."""
import hashlib
import html
import json
import os
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageOps, ImageSequence

from .config import IMAGES, OFFICE, TEXT

PREVIEW_VERSION = 1
DOCUMENT_CONVERSION_VERSION = 2


def command(args, timeout=180):
    result = subprocess.run(args, capture_output=True, timeout=timeout,
                            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    if result.returncode:
        raise ValueError(f'{Path(args[0]).stem} failed; check for corruption, password protection, or a missing audio track.')
    return result.stdout


def text_pdf(text, target):
    # Story paginates instead of truncating; CJK fallback fonts are embedded by MuPDF.
    body = '<html><body><pre>' + html.escape(text or '[Blank document]') + '</pre></body></html>'
    story = fitz.Story(html=body, user_css='body {font-family: sans-serif; font-size: 10pt;} pre {white-space: pre-wrap; overflow-wrap: break-word; font-family: sans-serif;}')
    bounds = fitz.Rect(0, 0, 595, 842)
    with fitz.DocumentWriter(str(target)) as writer:
        more = True
        while more:
            device = writer.begin_page(bounds)
            more, _ = story.place(bounds + (40, 40, -40, -40))
            story.draw(device)
            writer.end_page()


def html_pdf(text, target):
    """Render the actual document in an isolated browser, including inline charts."""
    from playwright.sync_api import sync_playwright, Error as BrowserError

    blocked = []
    with sync_playwright() as runtime:
        try:
            browser = runtime.chromium.launch(headless=True)
        except BrowserError:
            raise ValueError('HTML conversion requires Chromium. Run python -m playwright install chromium and retry.') from None
        try:
            context = browser.new_context(viewport={'width': 1000, 'height': 800},
                                          accept_downloads=False, service_workers='block')
            origin = 'https://folio.invalid/document'
            policy = ("default-src 'none'; script-src 'unsafe-inline'; style-src 'unsafe-inline'; "
                      "img-src data: blob:; font-src data:; connect-src 'none'; "
                      "frame-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'")
            served = False
            def route_request(route):
                nonlocal served
                if route.request.url == origin and not served:
                    served = True
                    route.fulfill(body=text, content_type='text/html; charset=utf-8',
                                  headers={'Content-Security-Policy': policy})
                else:
                    blocked.append(route.request.resource_type)
                    route.abort()
            context.route('**/*', route_request)
            context.route_web_socket('**/*', lambda socket: socket.close())
            page = context.new_page()
            page.add_init_script("""window.folioBlockedResources = 0;
                document.addEventListener('securitypolicyviolation', () => window.folioBlockedResources++);""")
            page_errors = []
            page.on('pageerror', lambda error: page_errors.append(True))
            page.emulate_media(media='screen', reduced_motion='reduce')
            page.goto(origin, wait_until='load', timeout=30000)
            page.wait_for_function("""async () => {
                document.querySelectorAll('details').forEach(element => element.open = true);
                document.querySelectorAll('img').forEach(image => image.loading = 'eager');
                document.querySelectorAll('body *').forEach(element => {
                    const style = getComputedStyle(element);
                    if (['fixed', 'sticky'].includes(style.position))
                        element.style.setProperty('position', 'static', 'important');
                    if (['auto', 'scroll'].includes(style.overflowY)) {
                        element.style.setProperty('height', 'auto', 'important');
                        element.style.setProperty('max-height', 'none', 'important');
                        element.style.setProperty('overflow', 'visible', 'important');
                    }
                });
                await document.fonts.ready;
                await Promise.all(Array.from(document.images, image => image.decode().catch(() => {})));
                await new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)));
                return true;
            }""", timeout=30000)
            if blocked or page.evaluate('window.folioBlockedResources > 0'):
                raise ValueError('HTML has blocked external resources. Embed images, styles, fonts, and scripts or provide a complete PDF. Incomplete pages were not sent to the model.')
            if page_errors or page.evaluate('Array.from(document.images).some(image => !image.naturalWidth)'):
                raise ValueError('HTML images or scripts failed to load. Fix the source or provide a complete PDF.')
            page.add_style_tag(content='''
                *, *::before, *::after { -webkit-print-color-adjust: exact !important; animation: none !important; transition: none !important; }
                img, svg, canvas, figure { break-inside: avoid; max-width: 100%; }
                pre { white-space: pre-wrap !important; overflow-wrap: break-word; }
            ''')
            page.pdf(path=str(target), format='A4', print_background=True,
                     margin={'top': '12mm', 'right': '12mm', 'bottom': '12mm', 'left': '12mm'})
        except BrowserError:
            raise ValueError('HTML rendering did not complete. Check images and scripts or provide a complete PDF.') from None
        finally:
            browser.close()


def office_supplement(source):
    """Expose content ordinary print/PDF export can leave out, as labeled extra pages."""
    ext = source.suffix.lower()
    if ext == '.xlsx':
        from openpyxl import load_workbook
        formulas = load_workbook(source, read_only=True, data_only=False)
        values = load_workbook(source, read_only=True, data_only=True)
        lines = ['Appendix: complete spreadsheet cells, including hidden sheets, rows, columns, and formulas',
                 'This appendix preserves non-empty cells. Consult the preceding PDF pages for layout and charts.']
        try:
            for sheet in formulas:
                lines.append(f'\nSheet: {sheet.title} (status: {sheet.sheet_state}）')
                for row, value_row in zip(sheet.iter_rows(), values[sheet.title].iter_rows()):
                    for cell, cached in zip(row, value_row):
                        if cell.value is not None:
                            value = str(cell.value)
                            if cell.data_type == 'f':
                                value += f' [cached value: {cached.value if cached.value is not None else "no cached value"}]'
                            lines.append(f'{cell.coordinate}: {value}')
        finally:
            formulas.close()
            values.close()
        return '\n'.join(lines)
    if ext not in {'.pptx', '.docx'}:
        return ''
    lines = []
    with zipfile.ZipFile(source) as archive:
        for name in sorted(archive.namelist()):
            include = (ext == '.pptx' and name.startswith('ppt/notesSlides/notesSlide') and name.endswith('.xml'))
            include |= name in {'word/comments.xml', 'word/footnotes.xml', 'word/endnotes.xml'}
            if include:
                root = ET.fromstring(archive.read(name))
                paragraphs = [''.join(n.text or '' for n in p.iter() if n.tag.endswith('}t'))
                              for p in root.iter() if p.tag.endswith('}p')]
                if any(paragraphs):
                    lines.extend([f'\nSource: {name}', *paragraphs])
    return '\n'.join(['Appendix: notes, comments, and footnotes (may repeat body content)', *lines]) if lines else ''


def convert_pdf(source, target):
    ext = source.suffix.lower()
    warnings = []
    if ext == '.pdf':
        shutil.copyfile(source, target)
    elif ext in IMAGES:
        with Image.open(source) as image, fitz.open() as doc:
            for frame in ImageSequence.Iterator(image):
                # Preserve all TIFF pages / GIF frames, with alpha composited on white.
                frame = ImageOps.exif_transpose(frame).convert('RGBA')
                bg = Image.new('RGBA', frame.size, 'white')
                bg.alpha_composite(frame)
                import io
                data = io.BytesIO()
                bg.convert('RGB').save(data, format='PNG')
                page = doc.new_page(width=frame.width * .75, height=frame.height * .75)
                page.insert_image(page.rect, stream=data.getvalue())
            doc.save(target)
    elif ext in TEXT:
        raw = source.read_bytes()
        if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
            text = raw.decode('utf-16')
        else:
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise ValueError('Text must be UTF-8 or UTF-16 with a BOM. Convert the encoding to avoid content loss.')
        if ext in {'.html', '.htm'}:
            html_pdf(text, target)
        else:
            text_pdf(text, target)
    elif ext in OFFICE:
        executable = shutil.which('soffice') or shutil.which('libreoffice')
        if not executable:
            raise ValueError('Install LibreOffice, add soffice to PATH, and retry.')
        with tempfile.TemporaryDirectory(prefix='folio-office-') as temp:
            temp = Path(temp)
            options = 'pdf'
            if ext in {'.ppt', '.pptx', '.odp'}:
                options = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'
            command([executable, f'-env:UserInstallation={(temp / "profile").as_uri()}',
                     '--headless', '--convert-to', options, '--outdir', str(temp), str(source)])
            result = temp / (source.stem + '.pdf')
            if not result.exists():
                raise ValueError('LibreOffice did not produce a PDF. The file may be protected or unsupported.')
            shutil.copyfile(result, target)
        supplement = office_supplement(source)
        if supplement:
            extra = target.with_name('supplement.pdf')
            text_pdf(supplement, extra)
            with fitz.open(target) as document, fitz.open(extra) as appendix:
                original_pages = len(document)
                document.insert_pdf(appendix)
                combined = target.with_name('combined.pdf')
                document.save(combined)
            combined.replace(target)
            warnings.append(f'Page {original_pages + 1} starts the appendix; page numbers refer to the converted PDF.')
        warnings.append('Review Office layout and fonts against the converted PDF. Embedded objects, animations, and hidden content may need manual verification.')
    else:
        raise ValueError(f'Unsupported: {ext or "no extension"} to PDF. Convert it to PDF separately, then import it again.')
    with fitz.open(target) as document:
        if document.needs_pass:
            raise ValueError('The PDF is password protected. Remove the password and import it again.')
        if len(document) == 0:
            raise ValueError('The PDF has no pages.')
        return len(document), warnings


def prepare_audio(source, folder):
    ffmpeg, ffprobe = shutil.which('ffmpeg'), shutil.which('ffprobe')
    if not ffmpeg or not ffprobe:
        raise ValueError('Install FFmpeg and ffprobe, add them to PATH, and retry.')
    # A complete recording is sent once. Compression changes format, never splits content.
    audio = folder / 'audio.mp3'
    command([ffmpeg, '-nostdin', '-y', '-i', str(source), '-vn', '-ac', '1', '-ar', '16000',
             '-c:a', 'libmp3lame', '-b:a', '48k', str(audio)], timeout=1800)
    if audio.stat().st_size >= 25_000_000:
        raise ValueError('The compressed full audio track still reaches the 25 MB API limit. It will not be split or truncated. Reduce the source size and import it again.')
    result = command([ffprobe, '-v', 'error', '-show_entries', 'format=duration', '-of', 'json', str(audio)])
    duration = float(json.loads(result)['format']['duration'])
    if duration <= 0:
        raise ValueError('The audio track has zero duration.')
    return audio, duration


def prepare_document(source, target):
    """Run in a process, since MuPDF must not be shared between Python threads."""
    count, warnings = convert_pdf(source, target)
    if count > 2000:
        raise ValueError('The document exceeds the 2,000-page local limit. It will not be split or truncated.')
    prepare_page_previews(target)
    return count, warnings


def prepare_page_previews(target):
    """Render local preview images in the converter process, including retries."""
    manifest = target.parent / 'page-previews.json'
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    if manifest.exists():
        cached = json.loads(manifest.read_text(encoding='utf-8'))
        if (cached.get('version') == PREVIEW_VERSION and cached.get('pdf_sha256') == digest
                and cached.get('page_count', 0) > 0
                and all((target.parent / 'pages' / f'{number}.png').exists()
                        for number in range(1, cached['page_count'] + 1))):
            return
    pages_dir = target.parent / 'pages'
    pages_dir.mkdir(exist_ok=True)
    with fitz.open(target) as document:
        count = len(document)
        if count > 2000:
            raise ValueError('The document exceeds the 2,000-page local limit. It will not be split or truncated.')
        for index, page in enumerate(document):
            # Preserve small labels in the local preview; these images are not uploaded.
            scale = min(3, 2400 / max(page.rect.width, page.rect.height))
            page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(pages_dir / f'{index + 1}.png')
    manifest.write_text(json.dumps({'version': PREVIEW_VERSION, 'pdf_sha256': digest, 'page_count': count},
                                   ensure_ascii=False, indent=2), encoding='utf-8')
